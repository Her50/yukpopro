"""
Génère la convention tripartite ABN / ACI / AFRI INSURANCE en .docx.
Toutes les références juridiques citées sont réelles (Code CIMA, lois camerounaises, OHADA).
"""

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


def set_cell_border(cell):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    for edge in ('top', 'left', 'bottom', 'right'):
        b = OxmlElement(f'w:{edge}')
        b.set(qn('w:val'), 'single')
        b.set(qn('w:sz'), '6')
        b.set(qn('w:color'), '000000')
        tcBorders.append(b)
    tcPr.append(tcBorders)


def add_heading(doc, text, level=1):
    h = doc.add_paragraph()
    h.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = h.add_run(text)
    run.bold = True
    if level == 0:
        run.font.size = Pt(16)
        run.font.color.rgb = RGBColor(0x0B, 0x2E, 0x4F)
        h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif level == 1:
        run.font.size = Pt(13)
        run.font.color.rgb = RGBColor(0x0B, 0x2E, 0x4F)
    else:
        run.font.size = Pt(11)
        run.font.color.rgb = RGBColor(0x14, 0x14, 0x14)
    h.paragraph_format.space_before = Pt(10)
    h.paragraph_format.space_after = Pt(4)
    return h


def add_para(doc, text, justify=True, bold=False, italic=False, size=11):
    p = doc.add_paragraph()
    if justify:
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style='List Bullet')
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(2)
    run = p.runs[0] if p.runs else p.add_run(text)
    if not p.runs:
        run = p.add_run(text)
    else:
        p.runs[0].text = text
        run = p.runs[0]
    run.font.size = Pt(11)
    return p


doc = Document()

# Marges
for section in doc.sections:
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)

# Style par défaut
style = doc.styles['Normal']
style.font.name = 'Calibri'
style.font.size = Pt(11)

# ========= EN-TÊTE =========
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = title.add_run("CONVENTION TRIPARTITE DE PARTENARIAT COMMERCIAL")
r.bold = True
r.font.size = Pt(16)
r.font.color.rgb = RGBColor(0x0B, 0x2E, 0x4F)

subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = subtitle.add_run("RELATIVE À LA SOUSCRIPTION D'ASSURANCES AUTOMOBILES\nDES VÉHICULES NEUFS COMMERCIALISÉS PAR LE CONCESSIONNAIRE")
r.bold = True
r.font.size = Pt(12)
r.font.color.rgb = RGBColor(0x33, 0x33, 0x33)

ref = doc.add_paragraph()
ref.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = ref.add_run("Référence : CTP/ABN-ACI-AFRI/2026/N°___")
r.italic = True
r.font.size = Pt(10)

doc.add_paragraph()

# ========= ENTRE LES SOUSSIGNÉS =========
add_heading(doc, "ENTRE LES SOUSSIGNÉS :", level=1)

add_para(doc,
    "La société ABN, société [forme juridique : SA / SARL] au capital social de "
    "[montant] FCFA, immatriculée au Registre du Commerce et du Crédit Mobilier de "
    "[ville] sous le numéro [RCCM N°…], dont le siège social est sis à [adresse], "
    "Cameroun, agissant en qualité de concessionnaire agréé pour la distribution et "
    "la commercialisation au Cameroun de véhicules neufs de marque chinoise "
    "[préciser marque(s)], représentée aux fins des présentes par "
    "[Nom, Prénoms, qualité du signataire], dûment habilité(e),"
)
add_para(doc, "Ci-après dénommée « le CONCESSIONNAIRE » ou « ABN »,", italic=True)
add_para(doc, "D'UNE PREMIÈRE PART,", bold=True)
doc.add_paragraph()

add_para(doc,
    "Le Cabinet ACI, société de courtage en assurance et réassurance, "
    "[forme juridique] au capital social de [montant] FCFA, immatriculée au RCCM "
    "de [ville] sous le numéro [RCCM N°…], titulaire de l'agrément en qualité de "
    "courtier d'assurance délivré par le Ministre des Finances de la République du "
    "Cameroun n° [N° d'agrément] du [date] conformément aux articles 501 et "
    "suivants du Code des assurances de la Conférence Interafricaine des Marchés "
    "d'Assurances (Code CIMA), dont le siège social est sis à [adresse], "
    "représentée par [Nom, Prénoms, qualité]."
)
add_para(doc, "Ci-après dénommée « le COURTIER » ou « ACI »,", italic=True)
add_para(doc, "D'UNE DEUXIÈME PART,", bold=True)
doc.add_paragraph()

add_para(doc,
    "La société AFRI INSURANCE, société anonyme d'assurance au capital social de "
    "[montant] FCFA, immatriculée au RCCM de [ville] sous le numéro [RCCM N°…], "
    "agréée pour pratiquer les opérations d'assurance non-vie, notamment la branche "
    "« Responsabilité civile véhicules terrestres à moteur » au sens de l'article "
    "328 du Code CIMA, par arrêté du Ministre des Finances n° [N°] du [date], "
    "dont le siège social est sis à [adresse], Cameroun, représentée par "
    "[Nom, Prénoms, qualité]."
)
add_para(doc, "Ci-après dénommée « l'ASSUREUR » ou « AFRI INSURANCE »,", italic=True)
add_para(doc, "D'UNE TROISIÈME PART,", bold=True)
doc.add_paragraph()

add_para(doc,
    "Ci-après désignées individuellement la « Partie » et collectivement les "
    "« Parties ».", italic=True
)

# ========= PRÉAMBULE =========
add_heading(doc, "PRÉAMBULE", level=1)

add_para(doc,
    "Considérant que toute personne physique ou morale, autre que l'État, dont la "
    "responsabilité civile peut être engagée en raison de dommages causés aux tiers "
    "par un véhicule terrestre à moteur, doit, pour faire circuler ledit véhicule, "
    "être couverte par une assurance garantissant cette responsabilité, "
    "conformément à l'article 200 du Code CIMA et à la Loi camerounaise n° 65/LF/9 "
    "du 22 mai 1965 portant assurance automobile obligatoire ;"
)
add_para(doc,
    "Considérant que la souscription effective de l'assurance de responsabilité "
    "civile constitue une condition légale préalable à la mise en circulation de "
    "tout véhicule neuf vendu par le CONCESSIONNAIRE sur le territoire de la "
    "République du Cameroun ;"
)
add_para(doc,
    "Considérant que le CONCESSIONNAIRE, dans le cadre de la commercialisation de "
    "véhicules neufs de marque chinoise, souhaite offrir à sa clientèle un service "
    "intégré comprenant, dès la livraison, la souscription d'un contrat "
    "d'assurance automobile conforme à la réglementation en vigueur ;"
)
add_para(doc,
    "Considérant que le COURTIER, intermédiaire d'assurance dûment agréé, "
    "représente le client auprès des entreprises d'assurance au sens de l'article "
    "501 du Code CIMA, et qu'il est tenu d'un devoir de conseil et d'information à "
    "l'égard de l'assuré ;"
)
add_para(doc,
    "Considérant que l'ASSUREUR dispose des agréments réglementaires requis pour "
    "souscrire les risques automobiles et qu'il accepte de mettre à disposition "
    "des clients du CONCESSIONNAIRE, par l'entremise du COURTIER, une offre "
    "d'assurance dédiée et tarifée dans des conditions préférentielles ;"
)
add_para(doc,
    "Considérant que les Parties entendent formaliser, par la présente Convention, "
    "un partenariat tripartite équilibré, transparent et conforme à la "
    "réglementation CIMA, à la législation camerounaise applicable et aux Actes "
    "uniformes de l'Organisation pour l'Harmonisation en Afrique du Droit des "
    "Affaires (OHADA), notamment l'Acte uniforme du 15 décembre 2010 portant sur "
    "le droit commercial général et l'Acte uniforme révisé du 30 janvier 2014 "
    "relatif au droit des sociétés commerciales et du groupement d'intérêt "
    "économique ;"
)

add_para(doc, "IL A ÉTÉ CONVENU ET ARRÊTÉ CE QUI SUIT :", bold=True)

# ============================================================
# TITRE I - DISPOSITIONS GÉNÉRALES
# ============================================================
add_heading(doc, "TITRE I — DISPOSITIONS GÉNÉRALES", level=1)

add_heading(doc, "Article 1 — Définitions", level=2)
add_para(doc, "Au sens de la présente Convention, les termes ci-après s'entendent comme suit :")
add_bullet(doc, "« Client » : toute personne physique ou morale acquéreur d'un véhicule neuf auprès du CONCESSIONNAIRE et qui souscrit, par l'entremise du COURTIER, un contrat d'assurance auprès de l'ASSUREUR.")
add_bullet(doc, "« Contrat d'assurance » : tout contrat d'assurance automobile souscrit dans le cadre de la présente Convention, comportant au minimum la garantie de responsabilité civile obligatoire prévue à l'article 200 du Code CIMA.")
add_bullet(doc, "« Prime » : la contrepartie financière du Contrat d'assurance versée par le Client à l'ASSUREUR par l'intermédiaire du COURTIER, dans les conditions de l'article 13 du Code CIMA.")
add_bullet(doc, "« Commission » : la rémunération due au COURTIER par l'ASSUREUR au titre de l'apport et de la gestion d'affaires, conformément à l'article 544 du Code CIMA.")
add_bullet(doc, "« Code CIMA » : le Code des assurances annexé au Traité instituant la Conférence Interafricaine des Marchés d'Assurances signé à Yaoundé le 10 juillet 1992, ensemble ses modifications subséquentes.")

add_heading(doc, "Article 2 — Objet", level=2)
add_para(doc,
    "La présente Convention a pour objet de définir le cadre, les modalités, les "
    "droits et obligations réciproques des Parties dans l'opération conjointe de "
    "promotion, de souscription, de gestion et de suivi des contrats d'assurance "
    "automobile relatifs aux véhicules neufs commercialisés par le CONCESSIONNAIRE "
    "au Cameroun."
)
add_para(doc,
    "Elle organise un partenariat exclusivement commercial entre des personnes "
    "morales juridiquement distinctes et indépendantes. Aucune des stipulations de "
    "la présente Convention ne saurait être interprétée comme créant entre les "
    "Parties une société, un groupement d'intérêt économique, une représentation "
    "salariée, un mandat de représentation générale ou tout autre lien qui "
    "porterait atteinte à l'indépendance juridique et opérationnelle de chaque "
    "Partie."
)

add_heading(doc, "Article 3 — Cadre juridique applicable", level=2)
add_para(doc,
    "Les Parties reconnaissent expressément que la présente Convention est régie, "
    "interprétée et exécutée conformément :"
)
add_bullet(doc, "au Traité CIMA du 10 juillet 1992 et au Code des assurances qui lui est annexé, notamment ses Livres I (Le contrat), II (Les assurances obligatoires), III (Les entreprises) et IV (Règles comptables) et ses Livres relatifs aux intermédiaires d'assurance ;")
add_bullet(doc, "à la Loi camerounaise n° 65/LF/9 du 22 mai 1965 portant assurance automobile obligatoire et ses textes d'application ;")
add_bullet(doc, "à la Loi camerounaise n° 2015/008 du 16 juillet 2015 portant création et fixant le régime indemnitaire du Fonds de garantie automobile ;")
add_bullet(doc, "à l'Acte uniforme OHADA du 15 décembre 2010 portant sur le droit commercial général ;")
add_bullet(doc, "à l'Acte uniforme OHADA révisé du 30 janvier 2014 relatif au droit des sociétés commerciales et du groupement d'intérêt économique ;")
add_bullet(doc, "aux usages professionnels du marché de l'assurance et du courtage, ainsi qu'aux décisions et circulaires de la Commission Régionale de Contrôle des Assurances (CRCA) et de la Direction Nationale des Assurances du Cameroun.")

add_heading(doc, "Article 4 — Durée et entrée en vigueur", level=2)
add_para(doc,
    "La présente Convention est conclue pour une durée d'un (1) an à compter de sa "
    "date de signature par les trois Parties. Elle est ensuite renouvelable par "
    "tacite reconduction par périodes successives d'un (1) an, sauf dénonciation "
    "notifiée par l'une quelconque des Parties aux deux autres par lettre "
    "recommandée avec accusé de réception ou par tout moyen écrit conférant date "
    "certaine, moyennant un préavis de trois (3) mois avant l'échéance."
)

# ============================================================
# TITRE II - OBLIGATIONS DU CONCESSIONNAIRE
# ============================================================
add_heading(doc, "TITRE II — OBLIGATIONS DU CONCESSIONNAIRE (ABN)", level=1)

add_heading(doc, "Article 5 — Promotion et orientation commerciale", level=2)
add_para(doc,
    "Le CONCESSIONNAIRE s'engage à promouvoir activement, auprès de toute personne "
    "acquéreur d'un véhicule neuf, la souscription d'un contrat d'assurance "
    "automobile par l'entremise du COURTIER auprès de l'ASSUREUR, en mettant "
    "notamment à disposition des espaces dédiés dans ses points de vente et en "
    "remettant les supports d'information validés conjointement par le COURTIER et "
    "l'ASSUREUR."
)
add_para(doc,
    "Le CONCESSIONNAIRE rappelle systématiquement à tout acquéreur le caractère "
    "obligatoire de l'assurance de responsabilité civile automobile en application "
    "de l'article 200 du Code CIMA et de la Loi n° 65/LF/9 du 22 mai 1965, sans "
    "toutefois imposer le choix de l'ASSUREUR, le client conservant son entière "
    "liberté contractuelle."
)

add_heading(doc, "Article 6 — Périmètre d'intervention et interdictions", level=2)
add_para(doc,
    "Le CONCESSIONNAIRE n'étant pas titulaire d'un agrément d'intermédiaire "
    "d'assurance au sens des articles 501 et suivants du Code CIMA, il s'interdit "
    "expressément :"
)
add_bullet(doc, "de présenter, conseiller ou souscrire pour le compte de tiers des opérations d'assurance autrement que par la simple mise en relation du Client avec le COURTIER ;")
add_bullet(doc, "d'encaisser des primes d'assurance, d'établir ou recevoir des chèques libellés à son ordre au titre de primes d'assurance, conformément à l'esprit des dispositions du Code CIMA encadrant l'encaissement des primes (notamment l'article 451) ;")
add_bullet(doc, "de modifier, traduire ou interpréter les conditions générales et particulières des contrats d'assurance ;")
add_bullet(doc, "d'utiliser les marques, logos et identifiants visuels du COURTIER ou de l'ASSUREUR sans autorisation écrite préalable.")

add_heading(doc, "Article 7 — Transmission des données techniques", level=2)
add_para(doc,
    "Pour chaque véhicule susceptible d'être assuré, le CONCESSIONNAIRE s'engage à "
    "transmettre au COURTIER, dans un délai de quarante-huit (48) heures à compter "
    "de la signature du bon de commande ou de la facture, les informations "
    "techniques et administratives suivantes : marque, modèle, type, numéro de "
    "châssis (VIN), puissance fiscale, date de première mise en circulation, "
    "valeur catalogue, ainsi que l'identité complète, l'adresse et les "
    "coordonnées de l'acquéreur."
)
add_para(doc,
    "Le CONCESSIONNAIRE garantit l'exactitude, la complétude et l'actualité de "
    "ces informations. Il ne saurait être tenu responsable des conséquences d'une "
    "déclaration de risque inexacte imputable au seul Client."
)

# ============================================================
# TITRE III - OBLIGATIONS DU COURTIER
# ============================================================
add_heading(doc, "TITRE III — OBLIGATIONS DU COURTIER (ACI)", level=1)

add_heading(doc, "Article 8 — Qualité, capacité et honorabilité", level=2)
add_para(doc,
    "Le COURTIER déclare et garantit qu'il satisfait, pendant toute la durée de "
    "la présente Convention, aux conditions de capacité professionnelle prévues "
    "par les articles 501 et suivants du Code CIMA et aux conditions "
    "d'honorabilité visées à l'article 506 du même Code. Il s'engage à informer "
    "sans délai les autres Parties de toute modification ou suspension affectant "
    "son agrément ou les conditions susvisées."
)
add_para(doc,
    "Le COURTIER intervient à titre exclusif comme représentant du Client auprès "
    "de l'ASSUREUR. Il ne représente en aucun cas l'ASSUREUR vis-à-vis des tiers."
)

add_heading(doc, "Article 9 — Devoir de conseil et d'information", level=2)
add_para(doc,
    "Conformément aux principes du Code CIMA et à la jurisprudence constante en "
    "matière de courtage, le COURTIER est tenu, à l'égard de chaque Client, d'un "
    "devoir d'information précontractuelle et d'un devoir de conseil. À ce titre, "
    "il :"
)
add_bullet(doc, "analyse les besoins du Client en fonction du véhicule acquis et de l'usage envisagé ;")
add_bullet(doc, "présente clairement les garanties obligatoires et les garanties facultatives complémentaires (vol, incendie, dommages tous accidents, bris de glace, individuelle conducteur, assistance, etc.) ;")
add_bullet(doc, "remet au Client une note d'information précontractuelle ainsi que les conditions générales du contrat ;")
add_bullet(doc, "recueille la proposition d'assurance dûment renseignée et signée par le Client.")

add_heading(doc, "Article 10 — Souscription et délivrance des attestations", level=2)
add_para(doc,
    "Le COURTIER procède, dans un délai maximum de soixante-douze (72) heures à "
    "compter de la réception de la proposition d'assurance complète et du "
    "paiement effectif de la prime, à la souscription du contrat auprès de "
    "l'ASSUREUR et à la remise au Client de l'attestation d'assurance dont la "
    "forme et le contenu sont fixés par les textes en vigueur."
)

add_heading(doc, "Article 11 — Encaissement et reversement des primes", level=2)
add_para(doc,
    "Conformément à l'article 13 du Code CIMA, dans sa rédaction issue du "
    "Règlement n° 0005/CIMA/PCMA/PCE/SG/2011 du 11 avril 2011 relatif à la "
    "souscription et au paiement de la prime, la prise d'effet du contrat est "
    "subordonnée au paiement de la prime par le Client. À défaut de paiement de "
    "la prime dans les conditions et délais prévus, le contrat est résilié de "
    "plein droit."
)
add_para(doc,
    "Le COURTIER s'engage à reverser à l'ASSUREUR les primes encaissées dans les "
    "conditions et délais convenus entre les Parties et conformes à la "
    "réglementation CIMA en vigueur. Aucun encaissement de prime ne saurait "
    "intervenir sur un compte autre qu'un compte bancaire dédié et identifié du "
    "COURTIER."
)

# ============================================================
# TITRE IV - OBLIGATIONS DE L'ASSUREUR
# ============================================================
add_heading(doc, "TITRE IV — OBLIGATIONS DE L'ASSUREUR (AFRI INSURANCE)", level=1)

add_heading(doc, "Article 12 — Souscription et émission des polices", level=2)
add_para(doc,
    "L'ASSUREUR s'engage à examiner et instruire toute proposition d'assurance "
    "transmise par le COURTIER dans le cadre de la présente Convention dans un "
    "délai n'excédant pas quarante-huit (48) heures ouvrées."
)
add_para(doc,
    "Les polices émises comportent l'ensemble des mentions obligatoires prévues "
    "par les dispositions du Code CIMA relatives aux mentions du contrat "
    "d'assurance (notamment article 8) et couvrent au minimum la responsabilité "
    "civile prévue à l'article 200 du Code CIMA."
)

add_heading(doc, "Article 13 — Tarification et conditions préférentielles", level=2)
add_para(doc,
    "L'ASSUREUR consent aux Clients orientés par le CONCESSIONNAIRE et présentés "
    "par le COURTIER une grille tarifaire dédiée, annexée à la présente "
    "Convention (Annexe 1), tenant compte du caractère neuf des véhicules "
    "concernés et du volume d'affaires attendu. Cette grille demeure conforme aux "
    "tarifs déposés et homologués auprès de l'autorité de contrôle compétente."
)

add_heading(doc, "Article 14 — Gestion des sinistres", level=2)
add_para(doc,
    "L'ASSUREUR s'engage à mettre en place un circuit accéléré de déclaration et "
    "de gestion des sinistres pour les Clients souscripteurs dans le cadre de la "
    "présente Convention. À cet effet :"
)
add_bullet(doc, "un interlocuteur dédié est désigné au sein du service Sinistres de l'ASSUREUR ;")
add_bullet(doc, "les délais d'expertise et de règlement sont conformes aux dispositions du Code CIMA, notamment celles relatives à l'indemnisation des victimes d'accidents de la circulation ;")
add_bullet(doc, "un reporting trimestriel sur la sinistralité du portefeuille est transmis au COURTIER et, sous forme agrégée et anonymisée, au CONCESSIONNAIRE.")

add_heading(doc, "Article 15 — Reversement des commissions", level=2)
add_para(doc,
    "Conformément à l'article 544 du Code CIMA, modifié par décision du Conseil "
    "des Ministres du 11 avril 2011, les commissions dues au COURTIER sont "
    "payées par l'ASSUREUR dans les trente (30) jours qui suivent la remise "
    "effective des primes à l'entreprise d'assurance. À défaut, les commissions "
    "non payées produisent de plein droit intérêt au double du taux d'escompte, "
    "dans la limite du taux d'usure, à compter de l'expiration du délai "
    "susvisé."
)
add_para(doc,
    "Les taux de commission applicables, fixés dans le respect des minima et "
    "maxima arrêtés par le Ministre en charge des assurances, sont précisés en "
    "Annexe 2."
)

# ============================================================
# TITRE V - DISPOSITIONS FINANCIÈRES ET COMMERCIALES
# ============================================================
add_heading(doc, "TITRE V — DISPOSITIONS FINANCIÈRES ET COMMERCIALES", level=1)

add_heading(doc, "Article 16 — Apport d'affaires et indemnité du concessionnaire", level=2)
add_para(doc,
    "En contrepartie de l'effort commercial déployé par le CONCESSIONNAIRE dans "
    "la promotion de l'offre d'assurance et dans la mise à disposition de ses "
    "infrastructures, le COURTIER lui rétrocède une indemnité d'apport d'affaires "
    "dont le taux et les modalités de calcul sont définis en Annexe 3."
)
add_para(doc,
    "Cette indemnité est exclusivement assise sur les commissions effectivement "
    "perçues par le COURTIER auprès de l'ASSUREUR. Elle ne saurait, en aucun cas, "
    "être qualifiée de commission d'intermédiation au sens du Code CIMA, le "
    "CONCESSIONNAIRE n'étant pas un intermédiaire d'assurance agréé. Elle "
    "constitue une rémunération de prestations commerciales et logistiques "
    "facturée par le CONCESSIONNAIRE au COURTIER, soumise au régime fiscal de "
    "droit commun applicable au Cameroun (TVA, IS, retenues à la source le cas "
    "échéant)."
)

add_heading(doc, "Article 17 — Reporting et reddition des comptes", level=2)
add_para(doc,
    "Les Parties conviennent de la tenue, à la fin de chaque trimestre civil, "
    "d'un comité de pilotage tripartite consacré à l'examen :"
)
add_bullet(doc, "du volume de contrats souscrits et des primes émises ;")
add_bullet(doc, "de l'état des encaissements et des reversements ;")
add_bullet(doc, "de l'évolution de la sinistralité et du ratio sinistres / primes (S/P) ;")
add_bullet(doc, "des commissions et indemnités d'apport versées au titre de la période ;")
add_bullet(doc, "des actions correctives et perspectives commerciales.")

# ============================================================
# TITRE VI - DISPOSITIONS DIVERSES
# ============================================================
add_heading(doc, "TITRE VI — DISPOSITIONS DIVERSES", level=1)

add_heading(doc, "Article 18 — Confidentialité et protection des données personnelles", level=2)
add_para(doc,
    "Les Parties s'engagent à préserver le caractère strictement confidentiel de "
    "toutes les informations échangées dans le cadre de la présente Convention. "
    "Cette obligation survit à l'expiration ou à la résiliation de la Convention "
    "pour une durée de cinq (5) ans."
)
add_para(doc,
    "Les Parties s'engagent à traiter les données à caractère personnel des "
    "Clients dans le strict respect de la Loi camerounaise n° 2010/012 du 21 "
    "décembre 2010 relative à la cybersécurité et à la cybercriminalité, ainsi "
    "que des dispositions sectorielles applicables. Chaque Partie est responsable "
    "des traitements qu'elle met en œuvre."
)

add_heading(doc, "Article 19 — Exclusivité", level=2)
add_para(doc,
    "Le CONCESSIONNAIRE accorde au COURTIER une exclusivité de présentation de "
    "l'offre d'assurance automobile au sein de ses points de vente pour la durée "
    "de la présente Convention. Cette exclusivité ne porte aucunement atteinte à "
    "la liberté de choix de l'assureur reconnue au Client par la réglementation. "
    "[À adapter selon la volonté réelle des parties — option non-exclusive "
    "possible.]"
)

add_heading(doc, "Article 20 — Force majeure", level=2)
add_para(doc,
    "Aucune des Parties ne pourra être tenue pour responsable de l'inexécution "
    "ou du retard dans l'exécution de l'une de ses obligations résultant d'un "
    "événement de force majeure tel que défini par la jurisprudence des "
    "juridictions camerounaises et de la Cour Commune de Justice et d'Arbitrage "
    "(CCJA) de l'OHADA."
)

add_heading(doc, "Article 21 — Résiliation", level=2)
add_para(doc,
    "Sans préjudice des cas de résiliation à l'échéance prévus à l'article 4, la "
    "présente Convention pourra être résiliée :"
)
add_bullet(doc, "de plein droit par toute Partie en cas de manquement grave de l'une des autres Parties à ses obligations contractuelles, demeuré non corrigé pendant un délai de trente (30) jours après mise en demeure restée infructueuse ;")
add_bullet(doc, "de plein droit en cas de retrait, suspension ou non-renouvellement de l'agrément du COURTIER ou de l'ASSUREUR par l'autorité compétente ;")
add_bullet(doc, "en cas d'ouverture d'une procédure collective d'apurement du passif au sens de l'Acte uniforme OHADA portant organisation des procédures collectives d'apurement du passif à l'encontre de l'une quelconque des Parties.")

add_para(doc,
    "La résiliation est sans effet sur l'exécution des contrats d'assurance "
    "souscrits antérieurement, qui se poursuivent jusqu'à leur terme conformément "
    "aux dispositions du Code CIMA."
)

add_heading(doc, "Article 22 — Règlement des différends", level=2)
add_para(doc,
    "En cas de différend né de l'interprétation, de l'exécution ou de la "
    "résiliation de la présente Convention, les Parties s'engagent à rechercher "
    "une solution amiable dans un délai de trente (30) jours à compter de la "
    "notification du différend par l'une d'elles."
)
add_para(doc,
    "À défaut d'accord amiable, le différend sera soumis, au choix de la Partie "
    "la plus diligente, soit aux juridictions camerounaises compétentes du "
    "ressort du siège social de l'ASSUREUR, soit à l'arbitrage de la Cour "
    "Commune de Justice et d'Arbitrage (CCJA) de l'OHADA conformément à son "
    "règlement d'arbitrage."
)

add_heading(doc, "Article 23 — Loi applicable", level=2)
add_para(doc,
    "La présente Convention est soumise au droit en vigueur en République du "
    "Cameroun et au droit communautaire CIMA et OHADA."
)

add_heading(doc, "Article 24 — Annexes", level=2)
add_para(doc, "Les annexes ci-après font partie intégrante de la présente Convention :")
add_bullet(doc, "Annexe 1 — Grille tarifaire dédiée des produits d'assurance automobile ;")
add_bullet(doc, "Annexe 2 — Taux et modalités des commissions versées au COURTIER ;")
add_bullet(doc, "Annexe 3 — Taux et modalités de l'indemnité d'apport d'affaires versée au CONCESSIONNAIRE ;")
add_bullet(doc, "Annexe 4 — Modèle de proposition d'assurance et liste des informations à transmettre.")

add_heading(doc, "Article 25 — Dispositions finales", level=2)
add_para(doc,
    "Toute modification de la présente Convention devra faire l'objet d'un "
    "avenant écrit signé par les trois Parties. La nullité éventuelle d'une "
    "stipulation n'entraînera pas la nullité des autres stipulations, qui "
    "demeureront pleinement applicables."
)

# ========= SIGNATURES =========
doc.add_paragraph()
add_para(doc,
    "Fait à ___________________, le ___ / ___ / 20___, en trois (3) exemplaires "
    "originaux, un pour chaque Partie.",
    bold=True
)
doc.add_paragraph()

# Tableau de signatures
sig_table = doc.add_table(rows=2, cols=3)
sig_table.autofit = True
headers = [
    "Pour le CONCESSIONNAIRE\n(ABN)",
    "Pour le COURTIER\n(ACI)",
    "Pour l'ASSUREUR\n(AFRI INSURANCE)",
]
for i, h in enumerate(headers):
    cell = sig_table.cell(0, i)
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(h)
    r.bold = True
    r.font.size = Pt(11)
    set_cell_border(cell)

sig_body = (
    "\n\n\n\nNom :\n\nQualité :\n\nSignature et cachet :"
)
for i in range(3):
    cell = sig_table.cell(1, i)
    cell.text = ""
    p = cell.paragraphs[0]
    r = p.add_run(sig_body)
    r.font.size = Pt(10)
    set_cell_border(cell)

# ========= NOTE DE BAS DE DOCUMENT =========
doc.add_paragraph()
note = doc.add_paragraph()
note.alignment = WD_ALIGN_PARAGRAPH.LEFT
r = note.add_run(
    "Note méthodologique — Sources juridiques utilisées :\n"
    "• Code CIMA, Livre I (Le contrat d'assurance) — articles 8 et 13 ;\n"
    "• Code CIMA, Livre II (Assurances obligatoires) — article 200 ;\n"
    "• Code CIMA, Livre V (Intermédiaires d'assurance) — articles 501, 506, 544 ;\n"
    "• Règlement n° 0005/CIMA/PCMA/PCE/SG/2011 du 11 avril 2011 (paiement de la prime) ;\n"
    "• Loi camerounaise n° 65/LF/9 du 22 mai 1965 portant assurance automobile obligatoire ;\n"
    "• Loi camerounaise n° 2015/008 du 16 juillet 2015 portant création du Fonds de garantie automobile ;\n"
    "• Loi camerounaise n° 2010/012 du 21 décembre 2010 relative à la cybersécurité ;\n"
    "• Acte uniforme OHADA du 15 décembre 2010 portant sur le droit commercial général ;\n"
    "• Acte uniforme OHADA du 30 janvier 2014 relatif aux sociétés commerciales et au GIE ;\n"
    "• Traité instituant la CIMA, Yaoundé, 10 juillet 1992."
)
r.italic = True
r.font.size = Pt(9)
r.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

output_path = r"c:\Users\23767\digitalisation_assurance\Convention_Tripartite_ABN_ACI_AFRI_INSURANCE.docx"
doc.save(output_path)
print(f"Document généré : {output_path}")
