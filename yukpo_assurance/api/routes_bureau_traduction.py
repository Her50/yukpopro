"""
Bureau Traduction — Traduction professionnelle avec terminologie africaine.
Routes : /api/v1/bureau/traduction/
"""
import logging
from pathlib import Path
from typing import Optional

from fastapi import APIRouter, Depends, HTTPException, File, Form, UploadFile
from fastapi.responses import Response
from pydantic import BaseModel

from core.auth import TokenData, get_current_user

logger = logging.getLogger("yukpo_assurance.api.bureau_traduction")
router = APIRouter()

_DATA_DIR = Path(__file__).parent.parent / "data" / "generated" / "bureau"
_DATA_DIR.mkdir(parents=True, exist_ok=True)

LANGUES = {
    "fr": "français", "en": "anglais", "es": "espagnol",
    "pt": "portugais", "ar": "arabe", "zh": "mandarin",
    "wo": "wolof", "ha": "haoussa", "sw": "swahili",
    "bm": "bambara", "yo": "yoruba", "ig": "igbo",
}

CONTEXTES = {
    "": "général",
    "comptabilite": "comptabilité / fiscalité SYSCOHADA",
    "juridique": "juridique / droit OHADA",
    "rh": "ressources humaines",
    "finance": "finance / banque / microfinance",
    "assurance": "assurance CIMA / FANAF",
    "ong": "ONG / développement international",
    "medical": "médical / santé publique",
    "technique": "ingénierie / BTP / mines",
    "agricole": "agriculture / agro-alimentaire",
    "commercial": "commercial / marketing",
}


class DemandeTraduction(BaseModel):
    contenu: str
    langue_source: str = "fr"
    langue_cible: str = "en"
    contexte_metier: Optional[str] = ""
    format_sortie: str = "texte"   # texte | docx


@router.post("/texte", summary="Traduction de texte avec terminologie métier africaine")
async def traduire_texte(
    demande: DemandeTraduction,
    current_user: TokenData = Depends(get_current_user),
):
    from core.ia_client import ia_client, ModeIA
    import time

    ls = LANGUES.get(demande.langue_source, demande.langue_source)
    lc = LANGUES.get(demande.langue_cible, demande.langue_cible)
    ctx = CONTEXTES.get(demande.contexte_metier or "", "général")

    prompt_sys = (
        f"Tu es un traducteur expert {ls} → {lc}, spécialisé en terminologie africaine francophone. "
        f"Contexte : {ctx}. "
        "Préserve : SYSCOHADA, OHADA, CEMAC, UEMOA, FCFA, COBAC, BCEAO, BEAC, noms propres, sigles. "
        "Retourne UNIQUEMENT la traduction, sans explication ni balises."
    )

    try:
        texte_traduit = await ia_client.appeler_ia(
            prompt_sys,
            f"Traduis ce texte :\n\n{demande.contenu}",
            mode=ModeIA.CLAUDE_RAPIDE,
        )
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Erreur IA : {e}")

    nb_mots_source = len(demande.contenu.split())
    nb_mots_cible = len(texte_traduit.split())

    fichier_id = None
    if demande.format_sortie == "docx":
        try:
            from docx import Document as _Doc
            doc = _Doc()
            doc.add_heading(f"Traduction {ls} → {lc}", 0)
            doc.add_paragraph(texte_traduit)
            buf = __import__("io").BytesIO()
            doc.save(buf)
            fichier_id = f"bureau_trad_{current_user.user_id}_{int(time.time())}.docx"
            (_DATA_DIR / fichier_id).write_bytes(buf.getvalue())
        except Exception:
            pass

    return {
        "texte_traduit": texte_traduit,
        "nb_mots_source": nb_mots_source,
        "nb_mots_cible": nb_mots_cible,
        "langue_source": demande.langue_source,
        "langue_cible": demande.langue_cible,
        "fichier_id": fichier_id,
    }


@router.post("/fichier", summary="Traduction d'un fichier (PDF, DOCX, TXT, image)")
async def traduire_fichier(
    fichier: UploadFile = File(...),
    langue_source: str = Form("fr"),
    langue_cible: str = Form("en"),
    contexte_metier: str = Form(""),
    current_user: TokenData = Depends(get_current_user),
):
    import time

    if fichier.size and fichier.size > 20 * 1024 * 1024:
        raise HTTPException(400, detail="Fichier trop volumineux (max 20 MB)")

    contenu_bytes = await fichier.read()
    ext = (fichier.filename or "").lower().rsplit(".", 1)[-1]

    # Extraction du texte selon le format
    texte_source = ""
    try:
        if ext in ("txt", "md", "csv"):
            texte_source = contenu_bytes.decode("utf-8", errors="ignore")

        elif ext in ("pdf",):
            try:
                import pypdf2 as PyPDF2
                reader = PyPDF2.PdfReader(__import__("io").BytesIO(contenu_bytes))
                texte_source = "\n".join(p.extract_text() or "" for p in reader.pages)
            except Exception:
                texte_source = contenu_bytes.decode("utf-8", errors="ignore")[:5000]

        elif ext in ("docx",):
            from docx import Document as _Doc
            doc = _Doc(__import__("io").BytesIO(contenu_bytes))
            texte_source = "\n".join(p.text for p in doc.paragraphs if p.text.strip())

        elif ext in ("png", "jpg", "jpeg", "webp"):
            import base64
            from core.ia_client import ia_client, ModeIA
            b64 = base64.b64encode(contenu_bytes).decode()
            mt = "image/png" if ext == "png" else "image/jpeg"
            texte_source = await ia_client.appeler_ia_vision(
                "Extrais tout le texte visible dans cette image. Retourne uniquement le texte extrait.",
                b64, mt, mode=ModeIA.CLAUDE_VISION,
            )
        else:
            texte_source = contenu_bytes.decode("utf-8", errors="ignore")[:8000]
    except Exception as e:
        raise HTTPException(400, detail=f"Impossible d'extraire le texte : {e}")

    if not texte_source.strip():
        raise HTTPException(400, detail="Aucun texte trouvé dans le fichier")

    # Traduction
    ls = LANGUES.get(langue_source, langue_source)
    lc = LANGUES.get(langue_cible, langue_cible)
    ctx = CONTEXTES.get(contexte_metier, "général")

    from core.ia_client import ia_client, ModeIA
    prompt_sys = (
        f"Tu es un traducteur expert {ls} → {lc}, spécialisé terminologie africaine. "
        f"Contexte : {ctx}. Préserve sigles, noms propres, FCFA, OHADA, SYSCOHADA. "
        "Retourne UNIQUEMENT la traduction."
    )
    try:
        texte_traduit = await ia_client.appeler_ia(
            prompt_sys,
            f"Traduis ce texte :\n\n{texte_source[:6000]}",
            mode=ModeIA.CLAUDE_RAPIDE,
        )
    except Exception as e:
        raise HTTPException(500, detail=f"Erreur traduction : {e}")

    # Export DOCX
    fichier_id = None
    try:
        from docx import Document as _Doc
        doc = _Doc()
        doc.add_heading(f"Traduction {ls} → {lc}", 0)
        doc.add_heading("Texte source", 1)
        doc.add_paragraph(texte_source[:2000] + ("…" if len(texte_source) > 2000 else ""))
        doc.add_heading("Texte traduit", 1)
        doc.add_paragraph(texte_traduit)
        buf = __import__("io").BytesIO()
        doc.save(buf)
        nom_base = (fichier.filename or "traduction").rsplit(".", 1)[0]
        fichier_id = f"bureau_trad_{current_user.user_id}_{int(time.time())}_{nom_base[:20]}.docx"
        (_DATA_DIR / fichier_id).write_bytes(buf.getvalue())
    except Exception:
        pass

    return {
        "texte_traduit": texte_traduit,
        "nb_mots_source": len(texte_source.split()),
        "nb_mots_cible": len(texte_traduit.split()),
        "langue_source": langue_source,
        "langue_cible": langue_cible,
        "fichier_id": fichier_id,
        "fichier_source": fichier.filename,
    }


@router.get("/fichier/{fichier_id}", summary="Télécharger fichier traduit")
async def telecharger_traduit(
    fichier_id: str,
    current_user: TokenData = Depends(get_current_user),
):
    if "/" in fichier_id or "\\" in fichier_id or ".." in fichier_id:
        raise HTTPException(400, detail="Nom de fichier invalide")
    chemin = _DATA_DIR / fichier_id
    if not chemin.exists():
        raise HTTPException(404, detail="Fichier introuvable")
    if f"_{current_user.user_id}_" not in fichier_id and current_user.role != "admin":
        raise HTTPException(403, detail="Accès refusé")
    ext = fichier_id.rsplit(".", 1)[-1].lower()
    mt = "application/vnd.openxmlformats-officedocument.wordprocessingml.document" if ext == "docx" else "application/octet-stream"
    return Response(content=chemin.read_bytes(), media_type=mt,
                    headers={"Content-Disposition": f'attachment; filename="{fichier_id}"'})
