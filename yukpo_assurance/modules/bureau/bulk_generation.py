"""
Sprint UX3 — Génération bulk depuis un CSV/XLSX uploadé.

Cas d'usage :
  - Banque uploade un CSV de 500 clients → génère 500 cartes de visite personnalisées
  - Hôtel uploade un CSV de 200 invités → 200 badges nominatifs
  - École uploade liste élèves → diplômes auto
  - Marketing : 1 000 visuels A/B-testés depuis variants CSV

Le LLM (Haiku) lit la 1ère ligne (headers) du CSV et propose un mapping :
  colonnes CSV → variables {placeholders} dans le brief template.

Puis pour chaque ligne, on génère un visuel en remplaçant les placeholders
et en déclenchant le pipeline standard (avec rate limiting concurrence=4 pour
ne pas saturer fal.ai/Replicate).

Limites :
  - Max 1000 lignes par CSV (au-delà : rejeter, suggérer split)
  - Tous les visuels même format/mode (l'utilisateur fixe une fois)
  - Coût total = nb_lignes × coût_unitaire (montré avant lancement)
"""
from __future__ import annotations

import asyncio
import csv
import io
import logging
from dataclasses import dataclass, field
from typing import Optional

logger = logging.getLogger("yukpo_assurance.bureau.bulk_generation")

MAX_LIGNES = 1000
MAX_CONCURRENCE = 4   # parallèles fal.ai/Replicate (éviter rate limit upstream)


@dataclass
class BulkRow:
    index: int
    data: dict
    statut: str = "pending"   # pending | running | done | failed
    download_url: Optional[str] = None
    erreur: Optional[str] = None
    cout_credits: float = 0.0


@dataclass
class BulkBatch:
    batch_id: str
    user_id: int
    compagnie_id: int
    total: int
    template_brief: str         # Ex: "Carte de visite pour {nom} {role} chez {entreprise}"
    cle_projet: str             # ex: "carte_visite"
    mode_visuel: str = "standard"
    rows: list[BulkRow] = field(default_factory=list)
    statut_global: str = "pending"   # pending | running | done | failed
    nb_done: int = 0
    nb_failed: int = 0
    cree_le: Optional[str] = None
    fini_le: Optional[str] = None


def parser_csv(contenu_bytes: bytes, encoding: str = "utf-8") -> list[dict]:
    """
    Parse un CSV en list[dict] (clé = header). Détecte automatiquement le
    délimiteur (`,`, `;`, `\\t`).
    Lève ValueError si plus de MAX_LIGNES.
    """
    try:
        texte = contenu_bytes.decode(encoding, errors="replace")
    except Exception as e:
        raise ValueError(f"CSV non décodable en {encoding} : {e}")

    # Détection délimiteur
    sample = texte[:4000]
    sniffer = csv.Sniffer()
    try:
        dialect = sniffer.sniff(sample, delimiters=",;\t|")
    except csv.Error:
        # fallback : virgule
        class _D(csv.Dialect):
            delimiter = ","; quotechar = '"'; doublequote = True
            skipinitialspace = True; lineterminator = "\n"; quoting = csv.QUOTE_MINIMAL
        dialect = _D

    reader = csv.DictReader(io.StringIO(texte), dialect=dialect)
    rows: list[dict] = []
    for i, r in enumerate(reader):
        if i >= MAX_LIGNES:
            raise ValueError(f"CSV trop grand (>{MAX_LIGNES} lignes). Splitter en plusieurs fichiers.")
        rows.append({k: (v or "").strip() for k, v in r.items() if k})
    return rows


def parser_xlsx(contenu_bytes: bytes) -> list[dict]:
    """Parse XLSX (1ère feuille). Utilise openpyxl."""
    try:
        import openpyxl
    except Exception as e:
        raise ValueError(f"openpyxl non dispo : {e}")
    wb = openpyxl.load_workbook(io.BytesIO(contenu_bytes), read_only=True, data_only=True)
    ws = wb.active
    rows_iter = ws.iter_rows(values_only=True)
    headers = [str(h or f"col_{i}").strip() for i, h in enumerate(next(rows_iter, ()) or [])]
    rows: list[dict] = []
    for i, row in enumerate(rows_iter):
        if i >= MAX_LIGNES:
            raise ValueError(f"XLSX trop grand (>{MAX_LIGNES} lignes).")
        rows.append({h: (str(v).strip() if v is not None else "") for h, v in zip(headers, row)})
    return rows


def detecter_format(filename: str, contenu_bytes: bytes) -> str:
    """Devine 'csv' ou 'xlsx' depuis le filename + magic bytes."""
    fl = (filename or "").lower()
    if fl.endswith(".csv") or fl.endswith(".tsv") or fl.endswith(".txt"):
        return "csv"
    if fl.endswith(".xlsx") or fl.endswith(".xlsm"):
        return "xlsx"
    # Magic bytes : XLSX = ZIP
    if contenu_bytes[:4] == b"PK\x03\x04":
        return "xlsx"
    return "csv"


async def proposer_mapping_llm(
    headers: list[str], sample_rows: list[dict],
    cle_projet_hint: Optional[str] = None,
) -> dict:
    """
    Demande à Haiku un mapping intelligent :
      - quelles colonnes correspondent à quels champs probables (nom/prénom/
        rôle/email/téléphone/entreprise/adresse/etc.)
      - quel template_brief proposer (avec {placeholders})
      - quel cle_projet recommander (carte_visite par défaut pour personnes)

    Output JSON :
      {
        "mapping": {"nom": "Nom", "role": "Fonction", ...},
        "template_brief": "Carte de visite pour {nom}, {role} chez {entreprise}.",
        "cle_projet_recommande": "carte_visite",
        "champs_obligatoires_manquants": [],
        "lignes_invalides_estimees": 0,
        "raison": "..."
      }
    """
    try:
        from core.ia_client import ia_client, ModeIA, ModelePrioritaire
        import json
        prompt = (
            f"Tu es ASSISTANT BULK GENERATION pour Yukpo Designer Pro.\n"
            f"L'utilisateur a uploadé un CSV/XLSX et veut générer N visuels personnalisés.\n\n"
            f"HEADERS détectés ({len(headers)} colonnes) :\n{headers}\n\n"
            f"3 PREMIÈRES LIGNES (échantillon) :\n"
            f"{json.dumps(sample_rows[:3], ensure_ascii=False, indent=2)}\n\n"
            f"{f'HINT cle_projet : {cle_projet_hint}' if cle_projet_hint else ''}\n\n"
            f"Tu dois retourner un JSON STRICT avec :\n"
            f"  - mapping     : {{champ_yukpo: nom_colonne_csv}} (ex: {{\"nom\": \"Nom\"}})\n"
            f"  - template_brief : phrase avec {{placeholders}} qui sera utilisée pour\n"
            f"    chaque ligne. Ex: \"Carte de visite pour {{nom}}, {{role}} chez {{entreprise}}.\n"
            f"    Email: {{email}}. Tél: {{telephone}}.\"\n"
            f"  - cle_projet_recommande (parmi: carte_visite, badge_nominatif, diplome,\n"
            f"    flyer_a5, faire_part_mariage, faire_part_bapteme, attestation, etiquette)\n"
            f"  - champs_obligatoires_manquants : list[str] (ex: [\"email\", \"telephone\"])\n"
            f"  - raison : 1 phrase justifiant le choix\n\n"
            f"Retourne UNIQUEMENT le JSON, sans markdown."
        )
        rep = await ia_client.appeler(
            prompt=prompt, mode=ModeIA.ANALYSE, json_attendu=True,
            forcer_modele=ModelePrioritaire.CLAUDE_HAIKU,
        )
        try:
            data = json.loads(rep.contenu)
        except json.JSONDecodeError:
            import re
            m = re.search(r'\{.*\}', rep.contenu, re.DOTALL)
            data = json.loads(m.group()) if m else {}
        data["_modele"] = rep.modele_utilise
        data["_tokens_in"] = rep.tokens_input
        data["_tokens_out"] = rep.tokens_output
        return data
    except Exception as e:
        logger.warning(f"[Bulk] mapping LLM échoué : {e}")
        return {
            "mapping": {h.lower(): h for h in headers},
            "template_brief": " ".join(f"{{{h.lower()}}}" for h in headers),
            "cle_projet_recommande": "carte_visite",
            "champs_obligatoires_manquants": [],
            "raison": "fallback heuristique (LLM indispo)",
        }


def remplir_template(template_brief: str, row_data: dict, mapping: dict) -> str:
    """
    Remplace les {placeholders} du template par les valeurs de la ligne.
    Mapping = {placeholder_yukpo: nom_colonne_csv}.
    """
    rendu = template_brief
    for placeholder, col_csv in mapping.items():
        valeur = row_data.get(col_csv, "")
        rendu = rendu.replace(f"{{{placeholder}}}", str(valeur))
    return rendu


async def estimer_cout_total(
    nb_lignes: int, mode_visuel: str, cle_projet: str,
) -> dict:
    """Estimation grossière du coût total (avant lancement)."""
    # Coût LLM par ligne (Haiku spec ~600 cr) + image générée selon mode
    cout_ia_par_ligne = 600
    cout_img_par_ligne = {
        "sans": 0, "standard": 15, "premium": 240,
        "ultra": 560, "ultra_plus": 1200,
    }.get(mode_visuel, 240)
    cout_creation = 20  # 1 page typique pour bulk
    total = (cout_ia_par_ligne + cout_img_par_ligne + cout_creation) * nb_lignes
    return {
        "nb_lignes": nb_lignes,
        "credits_par_ligne": cout_ia_par_ligne + cout_img_par_ligne + cout_creation,
        "credits_total": total,
        "fcfa_total": int(total * 0.6),   # 1 cr ≈ 0.6 FCFA user
        "duree_estimee_sec": nb_lignes * (3 if mode_visuel == "standard" else 8),
    }
