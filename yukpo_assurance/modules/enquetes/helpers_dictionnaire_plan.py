"""
Helpers complémentaires pour le module Enquêtes :
- Dictionnaire des variables (get/set/IA)
- Plan d'analyse (get/set/IA)
- Exports DOCX dédiés : questionnaire seul, plan d'analyse seul
"""
import base64
import io
import json
import re
from datetime import datetime, timezone

from modules.enquetes import gestionnaire_enquetes as ge


# ─── Dictionnaire des variables ────────────────────────────────────────────────

def get_dictionnaire_variables(etude_id: str) -> dict:
    etude = ge._etudes.get(etude_id)
    if not etude or not etude.formulaire:
        raise ValueError("Étude ou formulaire introuvable")
    form = etude.formulaire
    existing = getattr(form, "dictionnaire_variables", {}) or {}
    result: dict[str, dict] = {}
    for q in sorted(form.questions, key=lambda x: x.ordre):
        name = q.name_xlsform or ge._to_slug(q.libelle) or q.question_id
        prev = existing.get(name, {})
        result[name] = {
            "libelle": q.libelle,
            "type": q.type_question,
            "modalites": q.options,
            "obligatoire": q.obligatoire,
            "section": q.section_label,
            "description": prev.get("description", ""),
            "unite": prev.get("unite", ""),
            "categorie": prev.get("categorie", ""),
            "notes_metier": prev.get("notes_metier", ""),
            "modalites_detaillees": prev.get("modalites_detaillees", {}),
        }
    return {"etude_id": etude_id, "formulaire_id": form.formulaire_id, "variables": result}


def set_dictionnaire_variables(etude_id: str, variables: dict) -> dict:
    etude = ge._etudes.get(etude_id)
    if not etude or not etude.formulaire:
        raise ValueError("Étude ou formulaire introuvable")
    stored: dict[str, dict] = {}
    for name, meta in (variables or {}).items():
        stored[name] = {
            "description":          meta.get("description", ""),
            "unite":                meta.get("unite", ""),
            "categorie":            meta.get("categorie", ""),
            "notes_metier":         meta.get("notes_metier", ""),
            "modalites_detaillees": meta.get("modalites_detaillees", {}),
        }
    etude.formulaire.dictionnaire_variables = stored
    etude.updated_at = datetime.now(timezone.utc).isoformat()
    return get_dictionnaire_variables(etude_id)


async def generer_dictionnaire_ia(etude_id: str) -> dict:
    from core.ia_client import ia_client, ModeIA, ModelePrioritaire
    etude = ge._etudes.get(etude_id)
    if not etude or not etude.formulaire:
        raise ValueError("Étude ou formulaire introuvable")
    form = etude.formulaire
    qs = [
        {
            "name": q.name_xlsform or ge._to_slug(q.libelle) or q.question_id,
            "libelle": q.libelle,
            "type": q.type_question,
            "modalites": q.options,
            "section": q.section_label,
        }
        for q in sorted(form.questions, key=lambda x: x.ordre)
    ]
    prompt = (
        "Tu es un statisticien senior. À partir du contexte d'étude et de la liste des variables "
        "ci-dessous, rédige un dictionnaire des variables professionnel en JSON.\n\n"
        f"CONTEXTE :\nTitre : {etude.titre}\nTerrain : {etude.terrain}\n"
        f"Population : {etude.population_cible}\nQuestions de recherche : "
        + "\n".join(etude.questions_recherche) + "\n\n"
        f"VARIABLES DU FORMULAIRE :\n{json.dumps(qs, ensure_ascii=False, indent=2)}\n\n"
        'Retourne UNIQUEMENT un JSON de la forme : '
        '{"<name>": {"description":"...", "unite":"...", "categorie":"...", "notes_metier":"..."}}'
    )
    _reponse_ia = None
    try:
        rep = await ia_client.appeler(
            prompt=prompt, mode=ModeIA.ANALYSE,
            forcer_modele=ModelePrioritaire.CLAUDE_HAIKU,
        )
        _reponse_ia = rep
        raw = getattr(rep, "contenu", "") if hasattr(rep, "contenu") else str(rep)
        m = re.search(r"\{.*\}", raw, re.S)
        data = json.loads(m.group(0)) if m else {}
    except Exception:
        data = {}
    current = getattr(form, "dictionnaire_variables", {}) or {}
    for name, meta in data.items():
        if isinstance(meta, dict):
            current[name] = {**current.get(name, {}), **{k: v for k, v in meta.items() if v}}
    form.dictionnaire_variables = current
    etude.updated_at = datetime.now(timezone.utc).isoformat()
    result = get_dictionnaire_variables(etude_id)
    result["_reponse_ia"] = _reponse_ia
    return result


# ─── Plan d'analyse ───────────────────────────────────────────────────────────

def get_plan_analyse(etude_id: str) -> dict:
    etude = ge._etudes.get(etude_id)
    if not etude or not etude.formulaire:
        raise ValueError("Étude ou formulaire introuvable")
    return {
        "etude_id": etude_id,
        "plan_analyse": getattr(etude.formulaire, "plan_analyse", "") or "",
        "genere_automatiquement": bool(etude.analyse_quantitative),
    }


def set_plan_analyse(etude_id: str, plan_analyse: str) -> dict:
    etude = ge._etudes.get(etude_id)
    if not etude or not etude.formulaire:
        raise ValueError("Étude ou formulaire introuvable")
    etude.formulaire.plan_analyse = plan_analyse or ""
    etude.updated_at = datetime.now(timezone.utc).isoformat()
    return get_plan_analyse(etude_id)


async def generer_plan_analyse_ia(etude_id: str) -> dict:
    from core.ia_client import ia_client, ModeIA, ModelePrioritaire
    etude = ge._etudes.get(etude_id)
    if not etude:
        raise ValueError("Étude introuvable")
    if not etude.formulaire:
        raise ValueError("Aucun formulaire associé à cette étude")

    form = etude.formulaire
    variables = [
        {
            "name": q.name_xlsform or ge._to_slug(q.libelle) or q.question_id,
            "libelle": q.libelle, "type": q.type_question, "modalites": q.options[:8],
        }
        for q in sorted(form.questions, key=lambda x: x.ordre)
    ]
    qr = "\n".join("- " + q for q in etude.questions_recherche)
    prompt = (
        "Tu es un chercheur senior en méthodes mixtes. Rédige un PLAN D'ANALYSE professionnel "
        "(en markdown, structuré, détaillé) pour l'étude ci-dessous. "
        "Ce plan guidera l'analyse des données collectées.\n\n"
        f"CONTEXTE :\nTitre : {etude.titre}\nContexte : {etude.contexte}\n"
        f"Méthodologie : {etude.methodologie}\nMode : {etude.mode}\n"
        f"Terrain : {etude.terrain}\nPopulation : {etude.population_cible}\n"
        f"Questions de recherche :\n{qr}\n\n"
        f"VARIABLES COLLECTÉES ({len(variables)}) :\n"
        f"{json.dumps(variables, ensure_ascii=False, indent=2)}\n\n"
        "Structure attendue :\n"
        f"# Plan d'analyse — {etude.titre}\n\n"
        "## 1. Objectifs analytiques\n"
        "## 2. Préparation et nettoyage des données\n"
        "## 3. Analyses descriptives (par variable)\n"
        "## 4. Analyses inférentielles / croisées (hypothèses + tests)\n"
        "## 5. Analyse qualitative (si mixte/qualitatif)\n"
        "## 6. Hypothèses de recherche\n"
        "## 7. Livrables attendus\n\n"
        "Sois dense, précis, ancré dans le contexte. N'invente pas de variables."
    )
    try:
        rep = await ia_client.appeler(
            prompt=prompt, mode=ModeIA.REDACTION,
            forcer_modele=ModelePrioritaire.CLAUDE_OPUS,
        )
        texte = getattr(rep, "contenu", "") if hasattr(rep, "contenu") else str(rep)
    except Exception as e:
        raise ValueError(f"Erreur génération plan : {e}")

    form.plan_analyse = (texte or "").strip()
    etude.updated_at = datetime.now(timezone.utc).isoformat()
    result = get_plan_analyse(etude_id)
    result["_reponse_ia"] = rep
    return result


# ─── Exports DOCX dédiés ──────────────────────────────────────────────────────

def generer_questionnaire_docx_bytes(etude_id: str) -> bytes:
    """Génère un questionnaire format papier (version imprimable du formulaire)."""
    etude = ge._etudes.get(etude_id)
    if not etude or not etude.formulaire:
        raise ValueError("Étude ou formulaire introuvable")
    from docx import Document
    from docx.shared import Inches
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()
    titre = doc.add_heading(etude.formulaire.titre or etude.titre, 0)
    titre.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if etude.formulaire.description:
        pdesc = doc.add_paragraph(etude.formulaire.description)
        pdesc.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph("")
    entete = doc.add_paragraph()
    entete.add_run(
        "Date : _________________     Enquêteur : _________________     "
        "Code répondant : _________________"
    ).italic = True
    doc.add_paragraph("")

    current_section = None
    n = 0
    for q in sorted(etude.formulaire.questions, key=lambda x: x.ordre):
        if q.type_question in ("begin_group", "begin_repeat", "end_group", "end_repeat"):
            continue
        if q.section_label and q.section_label != current_section:
            current_section = q.section_label
            doc.add_heading(current_section, level=1)
        n += 1
        libelle = f"{n}. {q.libelle}" + (" *" if q.obligatoire else "")
        p = doc.add_paragraph()
        p.add_run(libelle).bold = True
        if q.hint:
            ph = doc.add_paragraph(q.hint)
            if ph.runs:
                ph.runs[0].italic = True
            ph.paragraph_format.left_indent = Inches(0.3)
        if q.type_question in ("select_one", "oui_non", "likert", "select_multiple"):
            for opt in q.options:
                doc.add_paragraph(f"☐  {opt}").paragraph_format.left_indent = Inches(0.3)
        elif q.type_question in ("rating", "range"):
            opts = q.options or ["1", "2", "3", "4", "5"]
            doc.add_paragraph("   ".join([f"☐ {o}" for o in opts])).paragraph_format.left_indent = Inches(0.3)
        elif q.type_question == "number":
            doc.add_paragraph("Réponse : ______________________").paragraph_format.left_indent = Inches(0.3)
        elif q.type_question == "date":
            doc.add_paragraph("Date : ____ / ____ / ________").paragraph_format.left_indent = Inches(0.3)
        elif q.type_question == "geopoint":
            doc.add_paragraph("Coordonnées GPS : lat ____________ lon ____________").paragraph_format.left_indent = Inches(0.3)
        elif q.type_question == "note":
            pn = doc.add_paragraph(q.libelle)
            pn.paragraph_format.left_indent = Inches(0.3)
        else:
            doc.add_paragraph("_" * 90).paragraph_format.left_indent = Inches(0.3)
            doc.add_paragraph("_" * 90).paragraph_format.left_indent = Inches(0.3)
        doc.add_paragraph("")

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.read()


def generer_plan_analyse_docx_bytes(etude_id: str) -> bytes:
    etude = ge._etudes.get(etude_id)
    if not etude or not etude.formulaire:
        raise ValueError("Étude ou formulaire introuvable")
    from docx import Document
    from docx.shared import Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()
    titre = doc.add_heading(f"Plan d'analyse — {etude.titre}", 0)
    titre.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph(
        f"Méthodologie : {etude.methodologie} | Terrain : {etude.terrain} | Mode : {etude.mode}"
    )
    doc.add_paragraph("")

    texte = getattr(etude.formulaire, "plan_analyse", "") or "_Plan d'analyse non généré._"
    for ligne in texte.split("\n"):
        ligne = ligne.rstrip()
        if not ligne:
            doc.add_paragraph("")
        elif ligne.startswith("# "):
            doc.add_heading(ligne.lstrip("# "), 1)
        elif ligne.startswith("## "):
            doc.add_heading(ligne.lstrip("# "), 2)
        elif ligne.startswith("### "):
            doc.add_heading(ligne.lstrip("# "), 3)
        else:
            p = doc.add_paragraph(ligne)
            p.paragraph_format.space_after = Pt(4)

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.read()
