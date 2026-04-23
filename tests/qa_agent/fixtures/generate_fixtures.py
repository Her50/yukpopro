"""Genere toutes les fixtures QA — idempotent, deterministe (seed fixe).

Usage:
    python tests/qa_agent/fixtures/generate_fixtures.py
"""
from __future__ import annotations

import asyncio
import logging
import random
import sys
from datetime import date, timedelta
from pathlib import Path

logging.basicConfig(level=logging.INFO, format="%(asctime)s [%(levelname)s] %(message)s")
log = logging.getLogger("fixtures")

SEED = 20260423
random.seed(SEED)

DIR = Path(__file__).resolve().parent
PAYS = ["Cameroun", "Côte d'Ivoire", "Sénégal", "Togo", "Bénin", "Burkina Faso", "Gabon"]
TYPES_SINISTRE = ["Auto matériel", "Auto corporel", "Incendie", "Vol", "Bris de glace",
                  "RC pro", "Dégâts des eaux"]
STATUTS = ["Ouvert", "En instruction", "Provisionné", "Réglé", "Clos sans suite"]


def _ensure_dir(p: Path) -> None:
    p.parent.mkdir(parents=True, exist_ok=True)


def gen_cv_pdf() -> Path:
    out = DIR / "cv_exemple.pdf"
    if out.exists():
        return out
    _ensure_dir(out)
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer

    doc = SimpleDocTemplate(str(out), pagesize=A4,
                            leftMargin=2 * cm, rightMargin=2 * cm,
                            topMargin=2 * cm, bottomMargin=2 * cm)
    s = getSampleStyleSheet()
    h1 = ParagraphStyle("h1", parent=s["Heading1"], fontSize=18, spaceAfter=12)
    body = s["BodyText"]

    elems = [
        Paragraph("MARIE NGONO MBALLA", h1),
        Paragraph("Actuaire confirmée — 8 ans d'expérience assurance vie & santé", body),
        Paragraph("Yaoundé, Cameroun · marie.ngono@example.cm · +237 6 99 12 34 56", body),
        Spacer(1, 12),
        Paragraph("<b>EXPERIENCE</b>", s["Heading2"]),
        Paragraph("<b>Actuaire Senior — ASSUR-CIMA Cameroun (2021-aujourd'hui)</b>", body),
        Paragraph("Tarification produits santé collective; provisionnement IBNR/CAT; "
                  "reporting CIMA trimestriel; modélisation DAV (10K assurés).", body),
        Spacer(1, 8),
        Paragraph("<b>Actuaire — NSIA Côte d'Ivoire (2018-2021)</b>", body),
        Paragraph("Pricing produits vie individuelle; étude rentabilité portefeuille; "
                  "appui projet Solvabilité régionale CIMA.", body),
        Spacer(1, 12),
        Paragraph("<b>FORMATION</b>", s["Heading2"]),
        Paragraph("ISFA Lyon — Master Actuariat (2017)", body),
        Paragraph("Université de Yaoundé I — Maîtrise Mathématiques Appliquées (2015)", body),
        Spacer(1, 12),
        Paragraph("<b>COMPETENCES</b>", s["Heading2"]),
        Paragraph("R, Python, SAS, Prophet, Solvabilité II, IFRS 17, normes CIMA.", body),
        Paragraph("Langues: Français (natif), Anglais (C1), Ewondo.", body),
    ]
    doc.build(elems)
    log.info(f"OK {out.name}")
    return out


def gen_protocole_enquete_pdf() -> Path:
    out = DIR / "protocole_enquete.pdf"
    if out.exists():
        return out
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak

    doc = SimpleDocTemplate(str(out), pagesize=A4,
                            leftMargin=2 * cm, rightMargin=2 * cm,
                            topMargin=2 * cm, bottomMargin=2 * cm)
    s = getSampleStyleSheet()
    body = s["BodyText"]

    elems = [
        Paragraph("PROTOCOLE D'ENQUETE", s["Title"]),
        Paragraph("Satisfaction clients — Assurance auto Douala", s["Heading2"]),
        Spacer(1, 12),
        Paragraph("<b>1. Contexte</b>", s["Heading2"]),
        Paragraph("La présente enquête vise à mesurer la satisfaction des clients "
                  "particuliers détenteurs d'un contrat assurance auto dans la ville "
                  "de Douala (Cameroun) sur l'année 2025. Le marché auto local connaît "
                  "une croissance de 7%/an mais souffre d'un NPS sectoriel faible (-12).", body),
        Spacer(1, 8),
        Paragraph("<b>2. Objectifs</b>", s["Heading2"]),
        Paragraph("- Mesurer la satisfaction globale (CSAT, NPS) par segment.<br/>"
                  "- Identifier les irritants principaux (déclaration sinistre, expertise, règlement).<br/>"
                  "- Évaluer la propension à la résiliation et à la recommandation.<br/>"
                  "- Capter les attentes en matière de digital (app, WhatsApp, USSD).", body),
        Spacer(1, 8),
        Paragraph("<b>3. Population cible</b>", s["Heading2"]),
        Paragraph("Particuliers, 25-60 ans, détenteurs d'un contrat auto >12 mois, "
                  "résidant dans Douala 1 à 5. Échantillon visé: 400 répondants, "
                  "stratifié par sexe/âge/type véhicule.", body),
        PageBreak(),
        Paragraph("<b>4. Méthodologie</b>", s["Heading2"]),
        Paragraph("Mixte: 320 questionnaires CAPI tablette (point de vente, agence) "
                  "+ 80 entretiens semi-directifs (durée 45min) sur dossiers sinistres "
                  "récents. Analyse quantitative SPSS + analyse qualitative thématique.", body),
        Spacer(1, 8),
        Paragraph("<b>5. Guide d'entretien — thèmes clés</b>", s["Heading2"]),
        Paragraph("Profil; choix de l'assureur; expérience souscription; "
                  "expérience sinistre; relation conseiller; canaux digitaux; "
                  "perception prix vs valeur; intention de renouvellement; verbatims.", body),
        Spacer(1, 8),
        Paragraph("<b>6. Variables sociodémographiques</b>", s["Heading2"]),
        Paragraph("Sexe, âge, CSP, niveau d'études, ancienneté du contrat, "
                  "type de véhicule, prime annuelle, nombre de sinistres déclarés.", body),
        PageBreak(),
        Paragraph("<b>7. Calendrier</b>", s["Heading2"]),
        Paragraph("Phase pilote: semaine S1; collecte: S2-S5; analyse: S6-S8; "
                  "restitution: S9.", body),
        Paragraph("<b>8. Confidentialité</b>", s["Heading2"]),
        Paragraph("Données anonymisées, conformes à la loi camerounaise n°2010/012 "
                  "sur la cybersécurité et la protection des données personnelles.", body),
        PageBreak(),
        Paragraph("<b>Annexe — exemples de questions</b>", s["Heading2"]),
        Paragraph("Q1. Quel est votre niveau global de satisfaction ? (1=très insatisfait, 5=très satisfait)<br/>"
                  "Q2. Recommanderiez-vous votre assureur ? (NPS 0-10)<br/>"
                  "Q3. Avez-vous déclaré un sinistre dans les 12 derniers mois ? (oui/non)<br/>"
                  "Q4. Si oui, délai de règlement perçu ? (texte libre)<br/>"
                  "Q5. Quel canal de contact préférez-vous ? (Agence, Tel, Email, WhatsApp, App)", body),
    ]
    doc.build(elems)
    log.info(f"OK {out.name}")
    return out


def gen_contrat_fr_docx() -> Path:
    out = DIR / "contrat_fr.docx"
    if out.exists():
        return out
    from docx import Document

    doc = Document()
    doc.add_heading("CONTRAT D'ASSURANCE MULTIRISQUE PROFESSIONNELLE", 0)
    doc.add_paragraph("Conditions générales et particulières — Référence n° MRP-2026-00428")

    doc.add_heading("Article 1 — Objet du contrat", level=1)
    doc.add_paragraph(
        "Le présent contrat a pour objet de garantir l'assuré contre les "
        "conséquences pécuniaires des risques couverts énumérés à l'article 4, "
        "dans les conditions et limites fixées aux conditions particulières."
    )
    doc.add_heading("Article 2 — Définitions", level=1)
    for term, definition in [
        ("Assuré", "personne physique ou morale dont les intérêts sont garantis"),
        ("Sinistre", "tout évènement de nature à mettre en jeu une garantie du contrat"),
        ("Franchise", "montant restant à la charge de l'assuré en cas de sinistre"),
        ("Prime", "somme due par l'assuré en contrepartie des garanties accordées"),
    ]:
        p = doc.add_paragraph()
        p.add_run(term + ": ").bold = True
        p.add_run(definition)

    doc.add_heading("Article 3 — Étendue territoriale", level=1)
    doc.add_paragraph(
        "Les garanties s'exercent sur le territoire du Cameroun, ainsi que dans "
        "tous les pays membres de la zone CIMA, sous réserve des dispositions "
        "particulières applicables au siège du sinistre."
    )

    doc.add_heading("Article 4 — Garanties accordées", level=1)
    table = doc.add_table(rows=5, cols=3)
    table.style = "Light Grid Accent 1"
    table.rows[0].cells[0].text = "Garantie"
    table.rows[0].cells[1].text = "Plafond (FCFA)"
    table.rows[0].cells[2].text = "Franchise (FCFA)"
    rows_data = [
        ("Incendie", "500 000 000", "500 000"),
        ("Dégâts des eaux", "150 000 000", "250 000"),
        ("Vol", "75 000 000", "1 000 000"),
        ("Responsabilité civile", "1 000 000 000", "0"),
    ]
    for i, (g, p, f) in enumerate(rows_data, 1):
        table.rows[i].cells[0].text = g
        table.rows[i].cells[1].text = p
        table.rows[i].cells[2].text = f

    doc.add_heading("Article 5 — Exclusions", level=1)
    doc.add_paragraph("Sont exclus des garanties :", style="List Bullet")
    for excl in [
        "les dommages résultant de la faute intentionnelle de l'assuré",
        "les dommages causés par des actes de terrorisme non couverts par la LAGT",
        "les dommages nucléaires de toute origine",
        "les dommages survenus en dehors de la zone CIMA",
    ]:
        doc.add_paragraph(excl, style="List Bullet")

    doc.add_heading("Article 6 — Procédure de déclaration de sinistre", level=1)
    doc.add_paragraph(
        "L'assuré dispose d'un délai de cinq (5) jours ouvrables à compter de la "
        "connaissance du sinistre pour le déclarer à l'assureur, par tout moyen "
        "écrit (courrier recommandé, courriel, formulaire digital). Pour le vol, "
        "ce délai est ramené à 48 heures."
    )

    doc.add_heading("Article 7 — Réassurance facultative", level=1)
    doc.add_paragraph(
        "Pour la part excédant le plafond de rétention de 250 000 000 FCFA, "
        "l'assureur cède en réassurance facultative auprès de réassureurs "
        "agréés sur la place de Lomé (CICA-RE) ou agréés CIMA."
    )

    doc.add_heading("Article 8 — Provision pour sinistres à payer", level=1)
    doc.add_paragraph(
        "L'assureur constitue une PSAP conforme aux normes prudentielles CIMA, "
        "calculée selon une méthode statistique éprouvée et révisée trimestriellement."
    )

    table2 = doc.add_table(rows=4, cols=2)
    table2.style = "Light Grid Accent 1"
    table2.rows[0].cells[0].text = "Indicateur"
    table2.rows[0].cells[1].text = "Valeur"
    table2.rows[1].cells[0].text = "Ratio combiné cible"
    table2.rows[1].cells[1].text = "< 95 %"
    table2.rows[2].cells[0].text = "Ratio sinistres"
    table2.rows[2].cells[1].text = "< 65 %"
    table2.rows[3].cells[0].text = "Ratio frais"
    table2.rows[3].cells[1].text = "< 30 %"

    doc.add_heading("Article 9 — Loi applicable et juridiction", level=1)
    doc.add_paragraph(
        "Le présent contrat est régi par le code des assurances CIMA. Tout "
        "litige sera porté devant le tribunal compétent de Yaoundé."
    )

    doc.save(str(out))
    log.info(f"OK {out.name}")
    return out


def gen_donnees_sinistres_xlsx(nom: str, n_lignes: int) -> Path:
    out = DIR / nom
    if out.exists():
        return out
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "sinistres"
    ws.append(["date", "type_sinistre", "montant_fcfa", "pays", "statut"])
    rng = random.Random(SEED)
    base = date(2024, 1, 1)
    for i in range(n_lignes):
        d = base + timedelta(days=rng.randint(0, 364))
        ws.append([
            d.isoformat(),
            rng.choice(TYPES_SINISTRE),
            rng.randint(50_000, 25_000_000),
            rng.choice(PAYS),
            rng.choice(STATUTS),
        ])
    wb.save(str(out))
    log.info(f"OK {out.name} ({n_lignes} lignes)")
    return out


def gen_portefeuille_clients_xlsx() -> Path:
    out = DIR / "portefeuille_clients.xlsx"
    if out.exists():
        return out
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "clients"
    ws.append(["client_id", "nom", "age", "pays", "prime_annuelle_fcfa",
               "nb_sinistres", "ltv_fcfa", "canal_acquisition"])
    rng = random.Random(SEED + 1)
    canaux = ["Agence", "Courtier", "Web", "App mobile", "WhatsApp", "Bancassurance"]
    prenoms = ["Marie", "Paul", "Aminata", "Jean", "Fatou", "Pierre", "Aïcha",
               "Samuel", "Khadidja", "Eric", "Awa", "Olivier"]
    noms = ["Mballa", "Kouassi", "Diop", "Nguema", "Touré", "Onana", "Sow",
            "Ngono", "Diallo", "Bekolo", "Ouattara", "Mensah"]
    for i in range(220):
        prime = rng.randint(80_000, 1_500_000)
        nb_sin = rng.choices([0, 1, 2, 3, 4, 5], weights=[55, 25, 12, 5, 2, 1])[0]
        ltv = int(prime * rng.uniform(2.5, 8.5) - nb_sin * rng.randint(150_000, 1_000_000))
        ws.append([
            f"CLI{i+1:05d}",
            f"{rng.choice(prenoms)} {rng.choice(noms)}",
            rng.randint(22, 75),
            rng.choice(PAYS),
            prime,
            nb_sin,
            max(0, ltv),
            rng.choice(canaux),
        ])
    wb.save(str(out))
    log.info(f"OK {out.name} (220 lignes)")
    return out


def gen_rapport_concurrent_pdf() -> Path:
    out = DIR / "rapport_concurrent.pdf"
    if out.exists():
        return out
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak

    doc = SimpleDocTemplate(str(out), pagesize=A4,
                            leftMargin=2 * cm, rightMargin=2 * cm,
                            topMargin=2 * cm, bottomMargin=2 * cm)
    s = getSampleStyleSheet()
    body = s["BodyText"]
    elems = [
        Paragraph("RAPPORT MARCHÉ — Assurance Auto Cameroun 2025", s["Title"]),
        Paragraph("Étude concurrentielle confidentielle", s["Heading2"]),
        Spacer(1, 12),
        Paragraph("<b>1. Synthèse</b>", s["Heading2"]),
        Paragraph("Le marché de l'assurance auto au Cameroun pèse 78 milliards FCFA "
                  "en 2025 (+7,2% YoY), porté par 4 acteurs principaux totalisant 71% "
                  "de parts de marché. La pénétration reste faible (12% du parc roulant).", body),
        PageBreak(),
        Paragraph("<b>2. Acteurs clés</b>", s["Heading2"]),
        Paragraph("ASSUR-CIMA: 24% PdM, leader sur le segment particuliers haut de gamme.<br/>"
                  "NSIA: 19% PdM, fort sur la flotte entreprise.<br/>"
                  "Allianz Cameroun: 16% PdM, expert produits internationaux.<br/>"
                  "Activa: 12% PdM, croissance digitale rapide (+45%).", body),
        PageBreak(),
        Paragraph("<b>3. Tendances</b>", s["Heading2"]),
        Paragraph("Digitalisation accélérée (souscription 100% mobile chez 2 acteurs); "
                  "pression réglementaire CIMA sur la solvabilité; "
                  "émergence du modèle pay-per-km (offre pilote chez Activa).", body),
        PageBreak(),
        Paragraph("<b>4. Risques</b>", s["Heading2"]),
        Paragraph("Sinistralité corporelle en hausse (+18%), tarifs sous pression, "
                  "concurrence accrue des courtiers digitaux régionaux.", body),
        PageBreak(),
        Paragraph("<b>5. Recommandations</b>", s["Heading2"]),
        Paragraph("Renforcer l'offre digitale (app mobile sinistres), nouer un partenariat "
                  "bancassurance, lancer un produit pay-per-km dès Q3 2026.", body),
    ]
    doc.build(elems)
    log.info(f"OK {out.name}")
    return out


def gen_notes_equipe_docx() -> Path:
    out = DIR / "notes_equipe.docx"
    if out.exists():
        return out
    from docx import Document
    doc = Document()
    doc.add_heading("Notes équipe — Comité produit auto Cameroun", 0)
    doc.add_paragraph("Date: 15 mars 2026 — Présents: Direction technique, Marketing, Souscription")
    doc.add_heading("Points discutés", level=1)
    for pt in [
        "Lancement offre auto connectée — pilote 200 clients à Yaoundé Q3.",
        "Refonte tarification jeunes conducteurs — bonus première année 15%.",
        "Partenariat bancassurance avec Afriland First Bank — signature prévue avril.",
        "Investissement IA traitement sinistres — ROI estimé 18 mois.",
        "Renégociation traités réassurance non-vie avec CICA-RE — fenêtre juin.",
    ]:
        doc.add_paragraph(pt, style="List Bullet")
    doc.add_heading("Décisions", level=1)
    doc.add_paragraph("- Validation budget pilote auto connecté: 120M FCFA.")
    doc.add_paragraph("- Recrutement 2 actuaires juniors avant fin Q2.")
    doc.add_paragraph("- Étude faisabilité produit pay-per-km à lancer immédiatement.")
    doc.save(str(out))
    log.info(f"OK {out.name}")
    return out


def gen_contrat_assurance_exemple_pdf() -> Path:
    out = DIR / "contrat_assurance_exemple.pdf"
    if out.exists():
        return out
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak

    doc = SimpleDocTemplate(str(out), pagesize=A4,
                            leftMargin=2 * cm, rightMargin=2 * cm,
                            topMargin=2 * cm, bottomMargin=2 * cm)
    s = getSampleStyleSheet()
    body = s["BodyText"]
    elems = [
        Paragraph("CONTRAT ASSURANCE HABITATION", s["Title"]),
        Paragraph("Référence: HAB-2026-007734 — Souscripteur: M. NDONGO Pierre", s["Heading2"]),
        Spacer(1, 12),
        Paragraph("<b>Clause 1 — Plafond global</b>", s["Heading2"]),
        Paragraph("Le plafond global d'indemnisation est fixé à 80 000 000 FCFA, "
                  "tous sinistres et toutes garanties confondues sur la durée du contrat. "
                  "<i>Cette clause limite drastiquement la protection de l'assuré pour "
                  "des sinistres multiples ou un sinistre majeur.</i>", body),
        PageBreak(),
        Paragraph("<b>Clause 2 — Délai de carence</b>", s["Heading2"]),
        Paragraph("Aucune indemnisation ne sera due pour les sinistres survenus "
                  "dans les 90 jours suivant la prise d'effet du contrat. "
                  "<i>Délai exceptionnellement long, défavorable à l'assuré.</i>", body),
        PageBreak(),
        Paragraph("<b>Clause 3 — Exclusions</b>", s["Heading2"]),
        Paragraph("Sont exclus: dommages causés par négligence de l'assuré, "
                  "dommages survenus en l'absence prolongée de l'assuré (>30 jours), "
                  "dommages d'origine non identifiée, dommages partiellement remboursés "
                  "par un tiers. <i>Liste très large, source de litiges fréquents.</i>", body),
        PageBreak(),
        Paragraph("<b>Clause 4 — Franchise</b>", s["Heading2"]),
        Paragraph("Franchise contractuelle: 15% du montant du sinistre, "
                  "minimum 500 000 FCFA. <i>Franchise proportionnelle élevée; "
                  "à comparer aux pratiques du marché (10%).</i>", body),
        PageBreak(),
        Paragraph("<b>Clause 5 — Résiliation unilatérale</b>", s["Heading2"]),
        Paragraph("L'assureur peut résilier le contrat sans motif moyennant "
                  "préavis de 30 jours. <i>Déséquilibre contractuel manifeste — "
                  "l'assuré ne dispose pas du même droit symétrique.</i>", body),
        PageBreak(),
        Paragraph("<b>Clause 6 — Indexation prime</b>", s["Heading2"]),
        Paragraph("La prime annuelle est révisable chaque échéance selon un indice "
                  "interne à l'assureur, sans plafond explicite. "
                  "<i>Risque de dérive tarifaire non encadrée.</i>", body),
    ]
    doc.build(elems)
    log.info(f"OK {out.name}")
    return out


def gen_document_fr_pdf() -> Path:
    out = DIR / "document_fr.pdf"
    if out.exists():
        return out
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak

    doc = SimpleDocTemplate(str(out), pagesize=A4,
                            leftMargin=2 * cm, rightMargin=2 * cm,
                            topMargin=2 * cm, bottomMargin=2 * cm)
    s = getSampleStyleSheet()
    body = s["BodyText"]
    elems = [
        Paragraph("Document technique — Réassurance facultative en zone CIMA", s["Title"]),
        Spacer(1, 12),
        Paragraph("La réassurance facultative est une opération par laquelle un assureur "
                  "cède un risque déterminé à un réassureur, qui décide librement de "
                  "l'accepter ou de le refuser. Elle s'oppose à la réassurance obligatoire "
                  "(par traité), où la cession est automatique.", body),
        PageBreak(),
        Paragraph("Dans la zone CIMA, les compagnies d'assurance ont l'obligation de céder "
                  "une part de leurs primes à la CICA-RE et aux réassureurs régionaux agréés, "
                  "afin de soutenir le développement du marché africain.", body),
        PageBreak(),
        Paragraph("La provision pour sinistres à payer (PSAP) doit être calculée selon une "
                  "méthode statistique reconnue. Le ratio combiné cible recommandé par les "
                  "régulateurs est inférieur à 95%.", body),
    ]
    doc.build(elems)
    log.info(f"OK {out.name}")
    return out


def gen_flyer_modele_png() -> Path:
    out = DIR / "flyer_modele.png"
    if out.exists():
        return out
    from PIL import Image, ImageDraw, ImageFont
    img = Image.new("RGB", (1024, 1024), (250, 245, 235))
    d = ImageDraw.Draw(img)
    try:
        font_big = ImageFont.truetype("arial.ttf", 80)
        font_med = ImageFont.truetype("arial.ttf", 36)
    except Exception:
        font_big = ImageFont.load_default()
        font_med = ImageFont.load_default()
    d.rectangle([0, 0, 1024, 200], fill=(20, 80, 140))
    d.text((40, 60), "ASSUR-CIMA", fill="white", font=font_big)
    d.text((40, 280), "Journée client", fill=(20, 80, 140), font=font_big)
    d.text((40, 380), "Samedi 12 avril 2026", fill=(80, 80, 80), font=font_med)
    d.text((40, 440), "Yaoundé — Hôtel Mont Fébé", fill=(80, 80, 80), font=font_med)
    d.rectangle([0, 900, 1024, 1024], fill=(220, 200, 80))
    d.text((40, 935), "Inscription gratuite — info@assur-cima.cm", fill=(20, 20, 20), font=font_med)
    img.save(str(out))
    log.info(f"OK {out.name}")
    return out


async def _tts_to_file(text: str, voice: str, mp3_path: Path) -> None:
    import edge_tts
    communicate = edge_tts.Communicate(text, voice)
    await communicate.save(str(mp3_path))


def gen_audio_reunion() -> Path:
    out = DIR / "reunion_60s_fr.m4a"
    txt_out = DIR / "reunion_60s_fr.txt"
    if out.exists() and txt_out.exists():
        return out
    transcript = (
        "Bonjour à tous, ouvrons le comité de souscription. "
        "Le dossier ACME concerne un risque industriel majeur à Douala. "
        "L'actuaire estime la prime à quatre-vingt-cinq millions FCFA. "
        "Le commercial demande un effort tarifaire de quinze pour cent. "
        "Le directeur technique recommande de céder cinquante pour cent en réassurance facultative. "
        "Décision: validation du dossier sous condition de cession à CICA-RE."
    )
    txt_out.write_text(transcript, encoding="utf-8")
    try:
        mp3 = DIR / "_reunion_tmp.mp3"
        asyncio.run(_tts_to_file(transcript, "fr-FR-DeniseNeural", mp3))
        try:
            from pydub import AudioSegment
            AudioSegment.from_mp3(str(mp3)).export(str(out), format="ipod")
            mp3.unlink(missing_ok=True)
        except Exception:
            mp3.rename(out.with_suffix(".mp3"))
            log.warning("ffmpeg/pydub absent: m4a non genere, mp3 conserve")
    except Exception as e:
        log.warning(f"edge-tts indisponible ({e}), audio skip")
    log.info(f"OK {out.name if out.exists() else txt_out.name}")
    return out


def gen_audio_terrain() -> Path:
    out = DIR / "audio_terrain_enquete.m4a"
    txt = DIR / "audio_terrain_enquete.txt"
    if txt.exists() and out.exists():
        return out
    transcript = (
        "Question: depuis combien de temps êtes-vous client chez nous ? "
        "Réponse: ça fait environ trois ans. "
        "Question: pouvez-vous nous parler de votre dernier sinistre ? "
        "Réponse: oui, j'ai eu un accrochage en janvier, le règlement a pris six semaines, "
        "je trouve ça un peu long. "
        "Question: quel canal préférez-vous pour nous joindre ? "
        "Réponse: WhatsApp clairement, j'utilise plus le téléphone classique."
    )
    txt.write_text(transcript, encoding="utf-8")
    try:
        mp3 = DIR / "_terrain_tmp.mp3"
        asyncio.run(_tts_to_file(transcript, "fr-FR-HenriNeural", mp3))
        try:
            from pydub import AudioSegment
            AudioSegment.from_mp3(str(mp3)).export(str(out), format="ipod")
            mp3.unlink(missing_ok=True)
        except Exception:
            mp3.rename(out.with_suffix(".mp3"))
    except Exception as e:
        log.warning(f"edge-tts indisponible ({e})")
    log.info(f"OK {out.name if out.exists() else txt.name}")
    return out


def main() -> int:
    DIR.mkdir(parents=True, exist_ok=True)
    log.info(f"Generation fixtures dans {DIR}")
    try:
        gen_cv_pdf()
        gen_protocole_enquete_pdf()
        gen_contrat_fr_docx()
        gen_donnees_sinistres_xlsx("donnees_sinistres.xlsx", 30)
        gen_donnees_sinistres_xlsx("sinistres_2024.xlsx", 50)
        gen_portefeuille_clients_xlsx()
        gen_rapport_concurrent_pdf()
        gen_notes_equipe_docx()
        gen_contrat_assurance_exemple_pdf()
        gen_document_fr_pdf()
        gen_flyer_modele_png()
        gen_audio_reunion()
        gen_audio_terrain()
    except Exception as e:
        log.error(f"echec generation: {e}", exc_info=True)
        return 1
    log.info("Fixtures generees avec succes")
    return 0


if __name__ == "__main__":
    sys.exit(main())
