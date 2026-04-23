"""Vérifications déterministes (pages, dimensions, structure, similarité)."""
from __future__ import annotations

from pathlib import Path
from typing import Any


def compter_pages_pdf(chemin: Path) -> int:
    from pypdf import PdfReader
    return len(PdfReader(str(chemin)).pages)


def dimensions_pdf_mm(chemin: Path, page: int = 0) -> tuple[float, float]:
    from pypdf import PdfReader
    p = PdfReader(str(chemin)).pages[page]
    box = p.mediabox
    pt_to_mm = 25.4 / 72.0
    return (float(box.width) * pt_to_mm, float(box.height) * pt_to_mm)


def texte_pdf(chemin: Path) -> str:
    from pypdf import PdfReader
    return "\n".join((p.extract_text() or "") for p in PdfReader(str(chemin)).pages)


def structure_docx(chemin: Path) -> dict[str, Any]:
    from docx import Document
    doc = Document(str(chemin))
    headings = [p.text for p in doc.paragraphs
                if p.style and p.style.name and p.style.name.startswith("Heading")]
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    return {
        "nb_paragraphs": len(paragraphs),
        "nb_tables": len(doc.tables),
        "nb_headings": len(headings),
        "headings": headings[:20],
        "nb_chars": sum(len(p) for p in paragraphs),
        "nb_images": len(doc.inline_shapes),
    }


def structure_pptx(chemin: Path) -> dict[str, Any]:
    from pptx import Presentation
    prs = Presentation(str(chemin))
    slides_data = []
    for slide in prs.slides:
        title = ""
        bullets = 0
        chars = 0
        for shape in slide.shapes:
            if not hasattr(shape, "text"):
                continue
            text = shape.text or ""
            chars += len(text)
            if shape.has_text_frame and not title:
                title = text.split("\n")[0]
            bullets += text.count("\n")
        slides_data.append({"title": title, "bullets": bullets, "chars": chars})
    return {"nb_slides": len(slides_data), "slides": slides_data}


def structure_xlsform(chemin: Path) -> dict[str, Any]:
    from openpyxl import load_workbook
    wb = load_workbook(str(chemin), data_only=True)
    sheets = wb.sheetnames
    survey = wb["survey"] if "survey" in sheets else None
    questions = []
    if survey:
        rows = list(survey.iter_rows(values_only=True))
        if rows:
            header = [str(h) for h in rows[0] if h is not None]
            type_idx = header.index("type") if "type" in header else 0
            for row in rows[1:]:
                if row and len(row) > type_idx and row[type_idx]:
                    questions.append(row[type_idx])
    return {
        "sheets": sheets,
        "has_survey": "survey" in sheets,
        "has_choices": "choices" in sheets,
        "has_settings": "settings" in sheets,
        "nb_questions": len(questions),
        "types_uniques": list({q for q in questions}),
    }


def image_dimensions(chemin: Path) -> tuple[int, int]:
    from PIL import Image
    with Image.open(chemin) as img:
        return img.size


def similarite_textes(a: str, b: str) -> float:
    from rapidfuzz import fuzz
    return fuzz.token_set_ratio(a, b) / 100.0


def detecter_langue(text: str) -> str:
    try:
        from lingua import Language, LanguageDetectorBuilder
        det = LanguageDetectorBuilder.from_languages(
            Language.FRENCH, Language.ENGLISH, Language.SPANISH, Language.PORTUGUESE
        ).build()
        lang = det.detect_language_of(text[:2000])
        return lang.iso_code_639_1.name if lang else "UNK"
    except Exception:
        return "UNK"


def chiffres_excel(chemin: Path, sheet: str | None = None) -> dict[str, Any]:
    from openpyxl import load_workbook
    wb = load_workbook(str(chemin), data_only=True)
    s = wb[sheet] if sheet else wb[wb.sheetnames[0]]
    rows = list(s.iter_rows(values_only=True))
    if not rows:
        return {}
    header = [str(h) if h else f"col{i}" for i, h in enumerate(rows[0])]
    data = rows[1:]
    aggregats: dict[str, Any] = {"nb_lignes": len(data), "colonnes": header}
    for i, col in enumerate(header):
        vals = [r[i] for r in data if i < len(r) and isinstance(r[i], (int, float))]
        if vals:
            aggregats[f"{col}_sum"] = sum(vals)
            aggregats[f"{col}_avg"] = sum(vals) / len(vals)
            aggregats[f"{col}_max"] = max(vals)
            aggregats[f"{col}_min"] = min(vals)
    return aggregats
