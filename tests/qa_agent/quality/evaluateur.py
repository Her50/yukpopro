"""Evaluation d'artefacts via Claude API.

Signature publique: evaluer_artefact(chemin, feature, attendus) -> dict
{ note: int 0-10, verdict: str, defauts: list[str], recommandations: list[str] }
"""
from __future__ import annotations

import base64
import json
import logging
import os
from pathlib import Path
from typing import Any

logger = logging.getLogger("qa_agent.quality")

try:
    from anthropic import Anthropic
except ImportError:
    Anthropic = None  # type: ignore


def _client() -> "Anthropic | None":
    if Anthropic is None:
        return None
    api_key = os.getenv("ANTHROPIC_API_KEY")
    if not api_key:
        return None
    return Anthropic(api_key=api_key)


def _extract_text(chemin: Path) -> str:
    suffix = chemin.suffix.lower()
    try:
        if suffix == ".pdf":
            from pypdf import PdfReader
            reader = PdfReader(str(chemin))
            return "\n\n".join((p.extract_text() or "") for p in reader.pages[:30])
        if suffix == ".docx":
            from docx import Document
            doc = Document(str(chemin))
            parts = [p.text for p in doc.paragraphs if p.text.strip()]
            for table in doc.tables:
                for row in table.rows:
                    parts.append(" | ".join(c.text for c in row.cells))
            return "\n".join(parts)
        if suffix == ".pptx":
            from pptx import Presentation
            prs = Presentation(str(chemin))
            slides = []
            for i, slide in enumerate(prs.slides, 1):
                texts = []
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text:
                        texts.append(shape.text)
                slides.append(f"--- Slide {i} ---\n" + "\n".join(texts))
            return "\n\n".join(slides)
        if suffix in (".xlsx", ".xls"):
            from openpyxl import load_workbook
            wb = load_workbook(str(chemin), data_only=True)
            parts = []
            for sheet in wb.sheetnames[:5]:
                ws = wb[sheet]
                parts.append(f"=== Sheet: {sheet} ===")
                for row in ws.iter_rows(values_only=True, max_row=50):
                    parts.append(" | ".join(str(c) if c is not None else "" for c in row))
            return "\n".join(parts)
        if suffix in (".txt", ".md", ".csv"):
            return chemin.read_text(encoding="utf-8", errors="ignore")
    except Exception as e:
        logger.warning(f"extraction texte echec {chemin.name}: {e}")
        return ""
    return ""


def _image_b64(chemin: Path) -> tuple[str, str] | None:
    suffix = chemin.suffix.lower().lstrip(".")
    media_type = {"png": "image/png", "jpg": "image/jpeg", "jpeg": "image/jpeg",
                  "webp": "image/webp", "gif": "image/gif"}.get(suffix)
    if not media_type:
        return None
    try:
        data = chemin.read_bytes()
        return media_type, base64.standard_b64encode(data).decode("ascii")
    except Exception:
        return None


_RUBRIQUE = """Tu es QA senior pour YukpoPro (SaaS assurance Afrique francophone, niveau pro international).

Feature: {feature}
Critères attendus: {attendus}

Évalue l'artefact ci-dessous selon ces axes (note globale /10):
1. Conformité au format/structure attendu
2. Qualité du contenu (profondeur, pertinence, absence d'hallucinations)
3. Niveau professionnel (typographie, lisibilité, sources citées si applicable)
4. Absence de placeholders / TODO / lorem ipsum / réponses LLM évasives
5. Adaptation contexte CIMA / OHADA / Afrique francophone si applicable

Renvoie STRICTEMENT un JSON valide (pas de markdown fencing) de la forme:
{{"note": <0-10>, "verdict": "PASS|PARTIAL|FAIL", "defauts": ["...", "..."], "recommandations": ["...", "..."]}}

PASS = note >= 7, PARTIAL = 5 <= note < 7, FAIL = note < 5."""


def _parse_json(text: str) -> dict[str, Any]:
    text = text.strip()
    if text.startswith("```"):
        text = text.strip("`").lstrip("json").strip()
    try:
        return json.loads(text)
    except Exception:
        start = text.find("{")
        end = text.rfind("}")
        if start >= 0 and end > start:
            try:
                return json.loads(text[start:end + 1])
            except Exception:
                pass
    return {"note": 0, "verdict": "FAIL",
            "defauts": ["JSON parse error", text[:200]],
            "recommandations": []}


def evaluer_artefact(chemin: str | Path, feature: str,
                     attendus: dict[str, Any] | None = None) -> dict[str, Any]:
    chemin = Path(chemin)
    attendus = attendus or {}

    if not chemin.exists():
        return {"note": 0, "verdict": "FAIL",
                "defauts": [f"fichier inexistant: {chemin}"],
                "recommandations": ["vérifier le téléchargement"]}

    cli = _client()
    if cli is None:
        logger.warning("ANTHROPIC_API_KEY absent, eval offline")
        return _eval_offline(chemin, feature, attendus)

    image = _image_b64(chemin)
    extrait = _extract_text(chemin)[:25000]

    user_prompt = _RUBRIQUE.format(
        feature=feature,
        attendus=json.dumps(attendus, ensure_ascii=False),
    )

    content: list[dict[str, Any]] = []
    if image:
        media_type, b64 = image
        content.append({
            "type": "image",
            "source": {"type": "base64", "media_type": media_type, "data": b64},
        })
    if extrait:
        content.append({"type": "text",
                        "text": f"--- EXTRAIT ARTEFACT ({chemin.name}) ---\n{extrait}"})
    elif not image:
        content.append({"type": "text",
                        "text": f"Aucun contenu extractible de {chemin.name} ({chemin.stat().st_size} octets)."})
    content.append({"type": "text", "text": user_prompt})

    try:
        resp = cli.messages.create(
            model=os.getenv("QA_CLAUDE_MODEL", "claude-opus-4-7"),
            max_tokens=1024,
            messages=[{"role": "user", "content": content}],
        )
        text = "".join(b.text for b in resp.content if hasattr(b, "text"))
        result = _parse_json(text)
        result["note"] = max(0, min(10, int(result.get("note", 0))))
        if "verdict" not in result:
            n = result["note"]
            result["verdict"] = "PASS" if n >= 7 else "PARTIAL" if n >= 5 else "FAIL"
        result.setdefault("defauts", [])
        result.setdefault("recommandations", [])
        return result
    except Exception as e:
        logger.error(f"Claude eval failed: {e}")
        return _eval_offline(chemin, feature, attendus)


def _eval_offline(chemin: Path, feature: str, attendus: dict[str, Any]) -> dict[str, Any]:
    """Fallback heuristique sans Claude."""
    if not chemin.exists() or chemin.stat().st_size < 100:
        return {"note": 1, "verdict": "FAIL",
                "defauts": ["fichier vide ou minuscule"],
                "recommandations": ["vérifier le pipeline de génération"]}
    text = _extract_text(chemin)
    defauts = []
    note = 6
    for marker in ("lorem ipsum", "TODO", "{{", "Je suis un LLM"):
        if marker.lower() in text.lower():
            defauts.append(f"placeholder détecté: {marker}")
            note -= 2
    if len(text) < 500 and chemin.suffix.lower() in (".pdf", ".docx"):
        defauts.append("contenu trop court (<500 chars)")
        note -= 2
    note = max(0, note)
    verdict = "PASS" if note >= 7 else "PARTIAL" if note >= 5 else "FAIL"
    return {"note": note, "verdict": verdict,
            "defauts": defauts,
            "recommandations": ["activer ANTHROPIC_API_KEY pour eval avancée"]}
